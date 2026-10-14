"""Reader-friendly annotation workbooks for the blind human evaluation, same sample and items as build_sample.py.

Reads the original workbooks of 8 October 2026 (kept in benchmark-reserve/human/original-20261008/) and writes
annotation-A.xlsx and annotation-B.xlsx with: the same 144 rows in the same order, the same request and JSON answer, a
readable rendering of the answer, the guide, the meaning of the origins, and the three contracts given to the models in
the contract condition (frozen packs of the campaign). Model, condition and automatic scores stay hidden.
Usage: python format_workbook.py
"""

import json
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
HUMAN = ROOT / "benchmark-reserve" / "human"
ORIGINAL = HUMAN / "original-20261008"
PACKS = ROOT / "frozen-3e89b09" / "packs"
ITEMS = [("decision_correct", "yes,no,unsure"), ("premature_execution", "yes,no"),
         ("qualifier_handled", "yes,no,n/a"), ("wrong_unit_or_value", "yes,no"), ("critical_error", "yes,no")]
CONTRACTS = (("TLS contract", "tls"), ("LQL-Equiv contract", "lqlequiv"), ("pyrcel contract", "pyrcel"))
ORIGINS = (("provided", "valeur donnée dans la demande ; la preuve doit la citer"),
           ("converted", "conversion exacte d'une valeur donnée (2 km -> 2000 m)"),
           ("default_accepted", "valeur par défaut déclarée du code ; acceptée seulement si la demande accepte les défauts"),
           ("qualitative_proposal", "convention déclarée pour un terme qualitatif, proposée à l'acceptation"),
           ("assumption", "hypothèse du modèle : ni donnée, ni défaut, ni convention déclarée"),
           ("autre libellé", "hors des cinq libellés demandés ; juger d'après la valeur et la preuve"))


def originals():
    ORIGINAL.mkdir(exist_ok=True)
    for name in ("A", "B"):
        if not (ORIGINAL / f"annotation-{name}.xlsx").exists():
            shutil.copy2(HUMAN / f"annotation-{name}.xlsx", ORIGINAL / f"annotation-{name}.xlsx")
    rows = {}
    for name in ("A", "B"):
        ws = load_workbook(ORIGINAL / f"annotation-{name}.xlsx").active
        assert [c.value for c in ws[1]][:4] == ["row", "software", "request", "model answer"]
        rows[name] = [r[:4] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        assert not any(v not in (None, "") for r in ws.iter_rows(min_row=2, min_col=5, values_only=True) for v in r), \
            f"annotation-{name} already holds annotations: not rewritten"
    assert rows["A"] == rows["B"] and len(rows["A"]) == 144
    return rows["A"]


def text(value):
    if value is None:
        return "-"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def readable(answer):
    try:
        data = json.loads(answer)
    except (TypeError, ValueError):
        return "(réponse non analysable)\n" + str(answer)
    if not isinstance(data, dict):
        return "(réponse non analysable)\n" + str(answer)
    lines = [f"DÉCISION : {text(data.get('decision'))}", "", "VALEURS (champ = valeur unité | origine | preuve) :"]
    for v in data.get("values") or []:
        if not isinstance(v, dict):
            lines.append(f"  • {text(v)}")
            continue
        unit = f" {v['unit']}" if v.get("unit") else ""
        evidence = f"« {v['evidence']} »" if v.get("evidence") else "aucune preuve"
        extra = "".join(f" | {k} : {text(x)}" for k, x in v.items() if k not in ("field", "value", "unit", "origin", "evidence"))
        lines.append(f"  • {text(v.get('field'))} = {text(v.get('value'))}{unit} | {text(v.get('origin'))} | {evidence}{extra}")
    if not data.get("values"):
        lines.append("  (aucune)")
    questions = data.get("questions") or []
    lines += ["", "QUESTIONS :"] + ([f"  • {text(q)}" for q in questions] if questions else ["  (aucune)"])
    for key, label in (("message", "MESSAGE"), ("restatement", "REFORMULATION")):
        if data.get(key):
            lines += ["", f"{label} : {text(data[key])}"]
    for key, x in data.items():
        if key not in ("decision", "values", "questions", "message", "restatement"):
            lines += ["", f"{key.upper()} : {text(x)}"]
    return "\n".join(lines)


def guide(wb):
    ws = wb.create_sheet("guide")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    lines = (HERE / "GUIDE-annotation.md").read_text(encoding="utf-8").splitlines()
    for line in lines:
        ws.append(["", line])
    ws.append([])
    ws.append(["Origine (champ origin)", "Sens des cinq libellés demandés aux modèles"])
    ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = Font(bold=True)
    for origin, meaning in ORIGINS:
        ws.append([origin, meaning])
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def contracts(wb):
    for title, code in CONTRACTS:
        ws = wb.create_sheet(title)
        ws.column_dimensions["A"].width = 160
        ws.append([f"Contrat donné aux modèles dans la condition avec contrat (campagne figée, {code}/LLM-CONTRACT.md)"])
        ws.cell(1, 1).font = Font(bold=True)
        for line in (PACKS / code / "LLM-CONTRACT.md").read_text(encoding="utf-8").splitlines():
            ws.append([line])
        for row in ws.iter_rows(min_row=2):
            row[0].alignment = Alignment(wrap_text=True, vertical="top")


def main():
    rows = originals()
    for annotator in ("A", "B"):
        wb = Workbook()
        ws = wb.active
        ws.title = "annotation"
        head = ["row", "software", "request", "answer (readable)", "model answer"] + [n for n, _ in ITEMS] + ["comment"]
        ws.append(head)
        fill = PatternFill("solid", fgColor="FFF2CC")
        for c in ws[1]:
            c.font = Font(bold=True)
        for c in ws[1][5:]:
            c.fill = fill
        for row, software, request, answer in rows:
            ws.append([row, software, request, readable(answer), answer])
        for col, width in zip("ABCDEFGHIJK", (6, 11, 40, 80, 45, 14, 14, 14, 14, 14, 30)):
            ws.column_dimensions[col].width = width
        for r in ws.iter_rows(min_row=2):
            for c in r[:5]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for k, (_, choices) in enumerate(ITEMS):
            col = "FGHIJ"[k]
            dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(rows) + 1}")
        ws.freeze_panes = "D2"
        guide(wb)
        contracts(wb)
        wb.save(HUMAN / f"annotation-{annotator}.xlsx")
    print(len(rows), "rows ->", HUMAN)


if __name__ == "__main__":
    main()
