"""Agreement for the blind human evaluation: Cohen's kappa and raw agreement, annotator A vs B and each annotator vs
the automatic scorer on the three items that have an automatic counterpart of the same definition; human-scored
contract effect (F010 vs F000) with Fisher's exact test.
wrong_unit_or_value has no automatic counterpart: the scorer's hallucination flag also counts defaults and conventions
used without acceptance, a different construct, so it is not compared.
Disagreements between A and B are listed in adjudication.xlsx (created once, never overwritten); the effect is reported
on adjudicated labels (agreed + adjudicated) when every disagreement is adjudicated, and separately on agreed labels.
With annotator B empty, results for A alone are provisional (no inter-annotator agreement).
Usage: python agreement.py   (reads benchmark-reserve/human/annotation-A.xlsx, annotation-B.xlsx, key.csv)
"""

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from scipy.stats import fisher_exact

HUMAN = Path(__file__).resolve().parents[3] / "benchmark-reserve" / "human"
ITEMS = ("decision_correct", "premature_execution", "qualifier_handled", "wrong_unit_or_value", "critical_error")
AUTO = {"decision_correct": "auto_decision_ok", "premature_execution": "auto_premature_execute",
        "qualifier_handled": "auto_qualifier_convention_ok"}
EFFECT = ("premature_execution", "critical_error")


def read(name):
    wb = load_workbook(HUMAN / f"annotation-{name}.xlsx")
    ws = wb["annotation"] if "annotation" in wb.sheetnames else wb.active
    head = [c.value for c in ws[1]]
    return {row[0]: dict(zip(head, row)) for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None}


def kappa(a, b):
    pairs = [(x, y) for x, y in zip(a, b) if x is not None and y is not None]
    if not pairs:
        return None, 0, None
    n = len(pairs)
    po = sum(x == y for x, y in pairs) / n
    ca, cb = Counter(x for x, _ in pairs), Counter(y for _, y in pairs)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return (None if pe == 1 else (po - pe) / (1 - pe)), n, po


def as_yesno(v):
    return None if v in ("", None) else ("yes" if float(v) > 0 else "no")


def adjudication(a, b, rows):
    """Final labels for the disagreements, from adjudication.xlsx; the file is created with an empty column if absent."""
    path = HUMAN / "adjudication.xlsx"
    disagreements = [(r, item, a[r][item], b[r][item]) for r in rows for item in ITEMS
                     if a[r][item] and b[r][item] and a[r][item] != b[r][item]]
    if not path.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "adjudication"
        ws.append(["row", "item", "A", "B", "final", "reason"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for d in disagreements:
            ws.append(list(d))
        wb.save(path)
        print(f"{len(disagreements)} disagreements written to {path.name}: fill 'final' together, keep A and B")
    ws = load_workbook(path)["adjudication"]
    final = {(r[0], r[1]): r[4] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None and r[4]}
    missing = [(r, item) for r, item, _, _ in disagreements if (r, item) not in final]
    return final, missing


def fisher(item, rows, label, key, note):
    table = {c: [0, 0] for c in ("F000", "F010")}
    for r in rows:
        v = label(r)
        if v in ("yes", "no"):
            table[key[r]["condition"]][0 if v == "yes" else 1] += 1
    p = fisher_exact([table["F010"], table["F000"]])[1]
    print(f"{item}: F010 yes/no={table['F010']}  F000 yes/no={table['F000']}  Fisher p={p:.3g} ({note})")


def main():
    a, b = read("A"), read("B")
    key = {int(r["row"]): r for r in csv.DictReader((HUMAN / "key.csv").open(encoding="utf-8"))}
    rows = sorted(set(a) & set(b))
    single = not any(b[r][item] for r in rows for item in ITEMS)  # second annotator not yet done: A alone
    print(f"{len(rows)} rows" + (" | PROVISIONAL: annotator B empty, A alone, no inter-annotator agreement" if single else ""))
    for item in ITEMS:
        line = f"{item:20}"
        if not single:
            k, n, po = kappa([a[r][item] for r in rows], [b[r][item] for r in rows])
            line += f" A vs B: kappa={k if k is None else round(k, 3)} agreement={po if po is None else round(po, 3)} n={n}"
        if item in AUTO:
            auto = [as_yesno(key[r][AUTO[item]]) for r in rows]
            for name, ann in (("A", a),) if single else (("A", a), ("B", b)):
                vals = [None if ann[r][item] in ("unsure", "n/a") else ann[r][item] for r in rows]
                k2, n2, po2 = kappa(vals, auto)
                line += f" | {name} vs scorer kappa={k2 if k2 is None else round(k2, 3)} agreement={po2 if po2 is None else round(po2, 3)} (n={n2})"
        else:
            line += " | no automatic counterpart"
        print(line)
    if single:
        for item in EFFECT:
            fisher(item, rows, lambda r: a[r][item], key, "annotator A alone, provisional")
        return
    final, missing = adjudication(a, b, rows)
    for item in EFFECT:
        fisher(item, rows, lambda r: a[r][item] if a[r][item] == b[r][item] else None, key, "agreed labels only")
        if any(i == item for _, i in missing):
            print(f"  adjudicated analysis pending: {sum(i == item for _, i in missing)} disagreements on {item} not adjudicated")
        else:
            fisher(item, rows, lambda r: a[r][item] if a[r][item] == b[r][item] else final.get((r, item)), key,
                   "adjudicated labels: agreed + adjudicated")


if __name__ == "__main__":
    main()
