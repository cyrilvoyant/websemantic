"""Blind human evaluation (evaluation/PREREG-toolrosella-and-human.md, part B): sample, annotation workbook, key.

Stratified random sample (seed 20261008): 3 models x 3 codes x {F000, F010} x 8 requests = 144 answers, among
requests with a single expected decision. Model, condition and automatic scores are hidden; rows are shuffled.
Outputs (private, in benchmark-reserve/human/): annotation-A.xlsx, annotation-B.xlsx (identical, one per annotator),
key.csv (row id -> answer file, model, condition, automatic scores).
"""

import csv
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
RESERVE = HERE.parents[2] / "benchmark-reserve"
E1 = HERE.parent / "e1"
OUT = RESERVE / "human"
MODELS = {"local_mistral-small-3.2-24b__20261008T085744Z-full": "mistral-small",
          "local_qwen2.5-72b__20261008T085744Z-full": "qwen-72b",
          "codestral-latest__final-20261008": "codestral"}
DOMAINS = ("tls", "lqlequiv", "pyrcel")
PER_CELL = 8
ITEMS = [("decision_correct", "yes,no,unsure"), ("premature_execution", "yes,no"),
         ("qualifier_handled", "yes,no,n/a"), ("wrong_unit_or_value", "yes,no"), ("critical_error", "yes,no")]


def cases():
    out = {}
    for d in DOMAINS:
        for corpus in ("pilot", "qualifiers"):
            for line in (RESERVE / d / f"{corpus}.jsonl").read_text(encoding="utf-8").splitlines():
                if line.strip():
                    c = json.loads(line)
                    if len(c.get("admissible", [c.get("decision")])) <= 1:  # single expected decision
                        out[c["id"]] = c
    return out


def main():
    rng = random.Random(20261008)
    cs = cases()
    sample = []
    for folder, model in MODELS.items():
        scores = {(r["case_id"], r["condition"], r["rep"]): r
                  for r in csv.DictReader((E1 / f"e1-scores-{folder}.csv").open(encoding="utf-8"))}
        for d in DOMAINS:
            ids = sorted(c for c in cs if cs[c]["domain"] == d)
            chosen = rng.sample(ids, PER_CELL)
            for cond in ("F000", "F010"):
                for cid in chosen:
                    rep = rng.choice((1, 2, 3))
                    f = RESERVE / "runs" / "e1" / folder / f"{cid}_{cond}_r{rep}.json"
                    if not f.exists():
                        continue
                    a = json.loads(f.read_text(encoding="utf-8"))
                    sample.append({"file": str(f.relative_to(RESERVE)), "model": model, "condition": cond, "case": cid,
                                   "domain": d, "request": a.get("request", ""), "answer": a.get("parsed") or a.get("raw_text"),
                                   "auto": scores.get((cid, cond, str(rep)), {})})
    rng.shuffle(sample)
    OUT.mkdir(parents=True, exist_ok=True)
    with (OUT / "key.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["row", "file", "model", "condition", "case", "domain", "auto_decision_ok", "auto_premature_execute",
                    "auto_qualifier_convention_ok", "auto_hallucination"])
        for i, s in enumerate(sample, 1):
            a = s["auto"]
            w.writerow([i, s["file"], s["model"], s["condition"], s["case"], s["domain"], a.get("decision_ok", ""),
                        a.get("premature_execute", ""), a.get("qualifier_convention_ok", ""), a.get("hallucination", "")])
    for annotator in ("A", "B"):
        wb = Workbook()
        ws = wb.active
        ws.title = "annotation"
        head = ["row", "software", "request", "model answer"] + [n for n, _ in ITEMS] + ["comment"]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, s in enumerate(sample, 1):
            ans = json.dumps(s["answer"], ensure_ascii=False, indent=1) if isinstance(s["answer"], dict) else str(s["answer"])
            ws.append([i, {"tls": "TLS", "lqlequiv": "LQL-Equiv", "pyrcel": "pyrcel"}[s["domain"]], s["request"], ans[:30000]])
        for col, width in zip("ABCDEFGHIJ", (6, 11, 50, 90, 14, 14, 14, 14, 14, 30)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for c in row[:4]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for k, (_, choices) in enumerate(ITEMS):
            col = "EFGHI"[k]
            dv = DataValidation(type="list", formula1=f'"{choices}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(sample) + 1}")
        ws.freeze_panes = "C2"
        wb.save(OUT / f"annotation-{annotator}.xlsx")
    print(len(sample), "answers ->", OUT)


if __name__ == "__main__":
    main()
